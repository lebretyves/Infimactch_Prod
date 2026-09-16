from pathlib import Path
import json,csv,re,zipfile
from datetime import datetime,timezone
from docx import Document
from docx.shared import Inches,Pt
root=Path(r'E:\Interimatch');out=root/'InfiMatch/docs/audits/v1-maquettes';data=json.loads((out/'audit-data.json').read_text(encoding='utf-8'));md=data['md'];rows=data['rows']
extra='''
## Défauts précis de données et de contrat
- INS-01 (P0) : Consentements transmet identité/email/ville à register mais pas les qualifications, la mobilité et les périodes. register élimine ensuite identité/ville du corps infirmier. Le succès crée un compte sans enregistrer le dossier saisi. Relier onboarding et profil, avec reprise après échec sans nouvelle création du compte.
- INS-02 (P0) : le RIB saisi n’est pas envoyé. Le backend /me/bank-details accepte uniquement un IBAN fictif contenant DEMO avec fictional:true ; il ne supporte pas BIC/titulaire. Brancher un IBAN réel échouerait. Aligner le parcours de démonstration et le contrat. Ne pas conserver ces données dans le stockage navigateur pour contourner ce manque.
- PRO-01 (P0) : compétences et expérience sont dans le DTO serveur, mais leur édition manque. Identité détaillée, ville lisible et diplôme/année nécessitent un mapping ou une évolution du modèle, pas seulement un appel API.
- CAN-01 (P0) : application.created_at attendu dans market.ts n’existe pas dans la table application, qui contient updated_at. La date affichée devient « Non précisée ». Définir une vraie date de dépôt ; ne pas présenter la dernière modification comme la date initiale.
- CAN-02 (P0) : le serveur écrit ACCEPTED après affectation ; le frontend traduit ASSIGNED mais pas ACCEPTED. Harmoniser sans mélanger les états candidature, mission et affectation.
- REC-01 (P1) : filtrage texte limité aux résultats déjà chargés ; il peut masquer l’existence d’une mission sur une autre page. Passer les critères au serveur et réinitialiser la pagination lors de leur changement.
- REC-02 (P1) : /listings/search classe par fraîcheur ; /me/matches classe selon le matching mais n’est pas consommé. L’écran actuel n’assure pas le classement métier demandé.
- ORG-01 (P1) : User frontend retient la première organisation mais perd son identifiant et son référent. Enrichir le contrat pour sauvegarder la bonne organisation et son contexte de rôle.
- ORG-02 (P0) : l’inscription entreprise visible impose FINESS et produit un établissement malgré le libellé établissement/agence. Prévoir le véritable parcours agence et ses affiliations autorisées. FINESS ne donne aucun droit d’accès.
- EXT-01 (P1) : qualité, provenance détaillée et critères inconnus/correspondance fournis par le backend ne sont pas affichés. Une annonce externe doit rester reliée au site source, sans score complet ni candidature interne inventés.
- FAV-01 (P1) : seuls les 50 premiers favoris sont lus ; au-delà, les étoiles peuvent ne plus refléter la collection. Prévoir pagination ou vérification ciblée.
## Navigation actuelle et navigation à compléter
**Intérimaire actuel :** accueil public → inscription → compte → accueil → missions → détail → candidater → candidatures. Profil et périodes sont partiellement éditables. Ruptures : dossier saisi non conservé, RPPS non saisissable, compétences non éditables, aucun écran favoris/historique/dossier ni accès confirmation.
**Établissement actuel :** connexion → accueil → missions → détail. Aucun examen/sélection/refus via interface. /profil possède une vue entreprise en lecture seule mais le menu l’exclut.
**Agence actuelle :** accueil et liste entreprise. Les API de création, publication, propositions et affectation n’ont pas leur parcours visuel.
**Routes proposées, pas encore existantes :** /dossier ; /favoris ; /historique ; /etablissement/missions/:id/candidatures ; /agence/missions/nouvelle ; /agence/missions/:id/modifier ; /agence/missions/:id/affectation. Réutiliser /profil avec menu adapté ou introduire une route organisation explicite. Une présentation minimale de l’établissement peut être intégrée aux favoris/détails.
Conserver critères/page de recherche et destination après connexion. Distinguer identifiants listing préfixés m_/e_ et UUID mission. Prévoir liens de retour, erreurs compréhensibles, page inconnue et rechargement des URL hébergées. Les profils proposés doivent rejoindre le processus de candidature/consentement requis avant affectation ; leur affichage ne suffit pas.
## Bilan par fonction V1
- F01 : classique raccordé ; onboarding global partiel et accès agence incomplet. Google demandé en complément reste non configuré.
- F02 / F02 bis : profils partiels ; RIB non raccordé et établissement non éditable.
- F03 : accueil/compteurs présents ; recommandations, résumés métier et accès manquants.
- F04 / F16 : règles déterministes, matching et explications backend présents ; intégration visuelle manquante. Aucun score ne remplace un prérequis bloquant.
- F05 : recherche réelle mais critères riches et présentation des résultats partiels.
- F06 : favoris mission partiels ; écran dédié/établissements/pagination manquants. Recherches sauvegardées hors V1.
- F07 : candidature/retrait présents ; examen établissement et suivi détaillé manquants.
- F08 : périodes éditables ; calendrier métier/mobilité intégrée à compléter.
- F09 : API historique affectations présente ; écran absent et timeline à enrichir.
- F12 : notifications lues ; trois automatismes prévus dans backend/workflows. Confirmation et liens métier non exploités ; recette complète à faire.
- F15 : acquisition publique réelle démontrée historiquement ; collecte bornée, retrait automatique des offres disparues encore à compléter.
- F17 / F19 : gestion missions et décision agence côté serveur ; écrans absents. Dossier candidat à enrichir pour la décision.
- F18 : vérification RPPS et retry côté serveur ; écran absent.
- F20 : annonces réelles/source visibles ; alertes qualité et informations manquantes non restituées.
- F22 : téléversement/lecture de document fictif côté serveur ; écran absent.
## Périmètre et divergences documentaires
Les maquettes 10 et 18 incluent des références professionnelles, exclues par le catalogue et le rectificatif actuels. Les retirer/annoter dans les maquettes ; ne pas les compter comme manque V1. Un RPPS absent bloque ; un fournisseur indisponible laisse en attente avec retry. Aucune validation manuelle agence de substitution.
RIB : présent dans F02 mais contrat actuel réservé à une démonstration fictive. Documents F22 : PDF/PNG/JPEG fictifs, 5 MiB maximum, JSON contentBase64/mime/fictional:true ; ne pas prévoir du multipart sans adaptation. GET RIB ne renvoie qu’une valeur masquée.
Hors V1 : références, attestation sur l’honneur, multi-CV, recherches sauvegardées, import CSV, signature électronique, contrats complets, pointage, paie, aide/chat avancés, duplication/export calendrier, gestion avancée des organisations. Les affiliations minimales nécessaires à l’agence restent requises.
Google est un complément demandé par l’utilisateur, pas une condition de la maquette de connexion classique. Bibliothèque et endpoints préparés, Client ID non créé, connexion réelle et Vercel non validés. La présence du code ne signifie pas que cela fonctionne sur Vercel.
Mot de passe oublié : route supplémentaire présente mais fonction toujours indisponible, sans API de réinitialisation. Implémenter le parcours ou retirer son entrée en attendant. Le minimum de 12 caractères, supérieur aux 8 de la maquette, est une règle serveur à conserver et expliquer, pas une fonctionnalité manquante.
Les dates/personnes des maquettes sont fictives et ne forment pas un état de base cohérent entre les pages. L’ancien INTERFACES.md doit être réconcilié avec le catalogue actuel.
## Exigences transverses : présent et restant à vérifier
- Architecture : React/TypeScript front et Node/TypeScript back, PostgreSQL, MongoDB, CLI et secrets présents. Les manques majeurs observés concernent les contrats de données et l’enchaînement des actions.
- Authentification/sécurité : session serveur, cookie HttpOnly, CSRF, hachage et permissions présents ; conserver ces protections dans chaque nouvel écran. Ce bilan ne constitue pas une revue de sécurité exhaustive.
- Automatismes : matching, confirmation après affectation manuelle, relance mission non pourvue. Vérifier déclenchements, reprise, absence de doublons et liens de destination après intégration. Exécution actuelle des workers non vérifiée ici.
- Accessibilité/responsive : preuve mobile limitée existante ; revue clavier/focus/erreurs/libellés/contrastes/tableaux sur tous les parcours à faire. Aucune conformité RGAA conclue.
- RGESN : documenter au moins deux pratiques et leurs preuves. Une pagination seule ne valide pas toutes les exigences.
- SEO : contrôler titres/meta publics, H1, URLs et sitemap ; non recetté exhaustivement.
- Données personnelles/exploitation : mentions présentes ; textes, rétention, droits, sauvegarde/restauration et configuration déployée à rapprocher des exigences et preuves. Aucune conclusion juridique globale.
- Hébergement : tests locaux ne prouvent pas Vercel. Vérifier URL API, cookies, CORS/CSRF, HTTPS, rechargement SPA et fournisseur Google sur l’environnement publié.
## Tests existants et recette manquante
Preuves disponibles, non relancées pour ce bilan : signup-regression.json (13 vérifications), integration-api.json (11 vérifications, 2026-09-15T20:32:09.355Z), profile-google.json (6 vérifications). Tests API frontend et compilation rapportés précédemment. Les anciens nombres de tests/couvertures backend sont des instantanés différents : ne pas les additionner ou les annoncer comme couverture actuelle des 18 vues.
La preuve integration-api précise que la candidature positive n’a pas été exercée car la fixture est RPPS_NOT_CHECKED. Elle prouve le refus correctement présenté, pas la chaîne de réussite. Les tests inscription vérifient compte/session/navigation, pas la persistance de tous les champs métier. Google est notamment testé désactivé, pas avec une identité fournisseur réelle.
### Recette à exécuter après corrections
- T01 : inscription complète, rechargement et comparaison des données /profile ; reprise après échec partiel sans compte doublon.
- T02 : compétences/expérience/RPPS ; cas trouvé, absent, indisponible/retry ; aucun contournement des blocages.
- T03 : agence crée, modifie et publie ; contrôle des rôles et affiliation établissement.
- T04 : filtres multiples, pagination, classement matching, inconnus externes, liens source et retour à la recherche.
- T05 : candidature admissible, double clic/rejeu, retrait, version modifiée et consentement renouvelé.
- T06 : établissement examine seulement ses candidats et sélectionne/refuse sans pouvoir affecter définitivement.
- T07 : agence affecte ; conflits refusés ; mission FILLED et candidature ACCEPTED correctement traduites.
- T08 : confirmation PENDING puis READY ; téléchargement autorisé, autre compte refusé, cas remplacé/annulé.
- T09 : favoris de trois types, retrait et plus de 50 entrées ; historique/calendrier cohérents avec affectations.
- T10 : document fictif valide, type/taille/accès invalides ; RIB fictif sauvegardé et masqué.
- T11 : trois workflows, reprises sans doublons, notifications et destination correcte.
- T12 : desktop/mobile/clavier, erreurs réseau/serveur, expiration session, URLs directes et recette hébergée ; Google séparément après configuration.
## Ordre recommandé
1. P0 : sauvegarde onboarding, profil complet et dossier RPPS/RIB fictif. Valider T01/T02/T10.
2. P0 : création/publication agence, examen établissement, affectation et PDF ; corriger dates/statuts. Valider T03/T05/T06/T07/T08.
3. P1 : recherche/matching expliqués, favoris, historique/calendrier, tableaux de bord, édition établissement. Valider T04/T09/T11.
4. P1/P2 : navigation, accueil, états d’erreur, accessibilité/exigences transverses et hébergement. Valider T12.
Clôture V1 : parcours positifs et négatifs validés avec données de démonstration maîtrisées, persistance après rechargement, respect des droits et vues métier atteignables. Aucun délai ni pourcentage global déduit du seul nombre de pages.
'''
md.extend(extra.strip().splitlines())
api=[('Auth','GET /auth/csrf ; POST /auth/register, /auth/login, /auth/logout ; GET /auth/me','Raccordé ; onboarding incomplet','auth/auth.module.ts'),('Google','GET /auth/google/config ; POST /auth/google/challenge ; POST /auth/google','Préparé, non configuré','auth/auth.module.ts'),('Profil','GET/PUT /profile','Partiel','profiles/profiles.module.ts'),('RPPS','PUT /profile/rpps ; POST /profile/rpps/retry','Non raccordé','profiles/profiles.module.ts'),('Recherche/détail','POST /listings/search ; GET /listings/:id','Partiel','listings/listings.module.ts'),('Public/correspondance','GET /listings/external ; GET /me/listings/:id/correspondence','Non raccordé directement','listings/listings.module.ts'),('Matching','GET /me/matches ; GET /missions/:id/candidates ; GET /matches/:id/explanation','Non raccordé ; respecter propriétaire des explications','matching/matching.module.ts'),('Favoris','GET/POST /me/favorites ; DELETE /me/favorites/:kind/:id','Partiel','listings/listings.module.ts'),('Référentiel','GET /facilities ; GET /reference-data/finess ; GET /reference-data/finess/:finess','Non raccordé','reference-data/finess.module.ts'),('Candidatures','POST /missions/:id/applications ; GET /me/applications ; POST /applications/:id/withdrawal','Raccordé partiellement ; réussite complète non recettée','missions/missions.module.ts'),('Examen','GET /missions/:id/applications ; POST /applications/:id/selection, /rejection','Non raccordé ; profil candidat à enrichir','missions/missions.module.ts'),('Missions','GET/POST /missions ; PUT /missions/:id ; POST /missions/:id/publish, /cancel, /reopen, /complete','Lecture seule raccordée','missions/missions.module.ts'),('Affectation','POST /missions/:id/assignments','Non raccordé','missions/missions.module.ts'),('Confirmation','GET /assignments/:id/confirmation ; GET /me/documents/:id','Non raccordé','organizations/organizations.module.ts'),('Historique','GET /me/history','Non raccordé ; timeline à enrichir','listings/listings.module.ts'),('Accueil/notifications','GET /dashboards ; GET /me/notifications ; POST /me/notifications/:id/read','Partiel','listings/listings.module.ts'),('Organisation','PUT /organizations/:id ; GET/POST /staffing-requests','Non raccordé','organizations/organizations.module.ts'),('Préférences','PUT /me/notification-preferences','Non raccordé ; infirmier','organizations/organizations.module.ts'),('Documents/RIB','GET/POST /me/documents ; GET /me/documents/:id ; GET/PUT /me/bank-details','Non raccordé ; données fictives','documents/documents.module.ts')]
with (out/'MATRICE_LIENS_API_V1.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f,delimiter=';');w.writerow(['Domaine','Routes relatives à /api/v1','Intégration','Preuve backend']);w.writerows(api)
md.append('## Annexe — contrats API et preuves')
md.append('Routes relatives à /api/v1. Les raccourcis de transitions reprennent le même préfixe de ressource. Inventaire fonctionnel ciblé, pas inventaire exhaustif des endpoints techniques.')
for a in api:md.append(f'- {a[0]} : {a[1]} — {a[2]}. Source : InfiMatch/backend/src/{a[3]}.')
anchors=[('infiMatch-front-end/src/services/auth.ts','export async function register'),('infiMatch-front-end/src/pages/inscription/Consentements.tsx','await register'),('infiMatch-front-end/src/services/market.ts','statusLabels'),('infiMatch-front-end/src/layouts/AppLayout.tsx','navigation.filter'),('infiMatch-front-end/src/router.tsx','export const router'),('InfiMatch/backend/src/database/schema.ts','CREATE TABLE application'),('InfiMatch/backend/src/documents/documents.module.ts','class BankDto')]
for file,needle in anchors:
 lines=(root/file).read_text(encoding='utf-8-sig').splitlines();n=next(i for i,l in enumerate(lines,1) if needle in l);md.append(f'- Preuve précise : {file}:{n} — {needle}.')
md.append('Pièces : sources.json (empreintes), MATRICE_PAGES_V1.csv (18 lignes), MATRICE_LIENS_API_V1.csv (19 groupes), planche-1/2/3.jpg (18 maquettes). backend-endpoints.json est un inventaire brut de décorateurs non normalisé ; utiliser la matrice API pour les chemins fonctionnels.')
content='\n\n'.join(md)+'\n';(out/'BILAN_V1_MAQUETTES_FRONT_BACK.md').write_text(content,encoding='utf-8')
doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
doc.sections[0].header.paragraphs[0].text='InfiMatch • Bilan fonctionnel V1 • Code local et maquettes'
for line in md:
 if line.startswith('# '):doc.add_heading(line[2:],0)
 elif line.startswith('## '):doc.add_heading(line[3:],1)
 elif line.startswith('### '):doc.add_heading(line[4:],2)
 elif line.strip():doc.add_paragraph(line.removeprefix('- ').replace('**',''),style='List Bullet' if line.startswith('- ') else None)
doc.add_heading('Annexe visuelle — maquettes fournies',1)
for i in range(1,4):
 doc.add_paragraph(f'Maquettes {(i-1)*6+1} à {i*6}')
 doc.add_picture(str(out/f'planche-{i}.jpg'),width=Inches(6.0))
doc.save(out/'BILAN_V1_MAQUETTES_FRONT_BACK.docx')
source=json.loads((out/'sources.json').read_text(encoding='utf-8-sig'))
assert len(rows)==18 and len({r[0] for r in rows})==18
assert sum(r[4]=='Absent' for r in rows)==6 and sum(r[4]=='Partiel' for r in rows)==11
assert all(s['hash_verified'] for s in source['screens'])
assert len(re.findall(r"path:\s*'",(root/'infiMatch-front-end/src/router.tsx').read_text(encoding='utf-8')))==21
for a in api:assert (root/'InfiMatch/backend/src'/a[3]).exists()
with zipfile.ZipFile(out/'BILAN_V1_MAQUETTES_FRONT_BACK.docx') as z:assert z.testzip() is None
proof={'generated_at':datetime.now(timezone.utc).isoformat(),'pages':18,'absent':6,'partial':11,'covered_drawn_scope':1,'api_groups':len(api),'front_paths':21,'hashes_verified':True,'docx_valid':True,'new_application_tests':False}
(out/'validation-bilan.json').write_text(json.dumps(proof,indent=2),encoding='utf-8');print(json.dumps(proof))
