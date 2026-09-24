from pathlib import Path
import re,sys
from docx import Document
from docx.shared import Pt
ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/('docs_intern' if (ROOT/'docs_intern/REQUIREMENTS_V1.md').exists() else 'docs')
PAIRS={'CATALOGUE_FONCTIONNALITES_V1_ACTUALISE.docx':'REQUIREMENTS_V1.md','CONTENU_VAULT_V1.docx':'VAULT_V1.md','RECTIFICATIF_CATALOGUE_V1.docx':'RECTIFICATIF_CATALOGUE_V1.md','EXPLICATION_MATCHING_DONNEES_MANQUANTES.docx':'EXPLICATION_MATCHING_DONNEES_MANQUANTES.md'}
def plain(text):
 text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',text)
 return text.replace('**','').replace('`','')
def render(vault_only=False):
 for dest,source in PAIRS.items():
  if vault_only and source!='VAULT_V1.md':continue
  d=Document();d.sections[0].header.paragraphs[0].text='InfiMatch | Documentation du projet'
  d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(11)
  d.add_paragraph('Révision documentaire du 24 septembre 2026. État de livraison : docs/ETAT_COURANT.md. Source maintenue : '+source+'. Les preuves datées et les réserves conservent leur portée ; ce document ne constitue pas une recette de production.')
  inside=False;table=None
  for line in (DOCS/source).read_text(encoding='utf-8').splitlines():
   if line.startswith('```'):
    inside=not inside;table=None
    if inside:d.add_paragraph('Schéma ou code : consulter la version Markdown liée ci-dessus.')
    continue
   if inside:continue
   if line.startswith('|'):
    cells=[x.strip() for x in line.strip('|').split('|')]
    if all(set(x)<=set('-: ') for x in cells):continue
    if table is None:table=d.add_table(rows=0,cols=len(cells));table.style='Light Shading Accent 1'
    for cell,value in zip(table.add_row().cells,cells):cell.text=plain(value)
   else:
    table=None
    if line.startswith('#'):d.add_heading(plain(line.lstrip('# ').strip()),min(3,len(line)-len(line.lstrip('#'))))
    elif line.strip():d.add_paragraph(plain(line[2:] if line.startswith('- ') else line),style='List Bullet' if line.startswith('- ') else None)
  d.sections[0].footer.paragraphs[0].text='Projet étudiant | Référence courante : fichier Markdown du dépôt'
  d.save(DOCS/dest)
  print('Generated',dest)
if __name__=='__main__':render('--vault-only' in sys.argv)
