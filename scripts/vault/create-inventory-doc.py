from pathlib import Path
from datetime import datetime
import json, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(r'E:\Interimatch\InfiMatch')
doc=Document()
sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
normal=doc.styles['Normal']
normal.font.name='Calibri'
normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(6)
for name in ['Title','Heading 1','Heading 2']:
 doc.styles[name].font.color.rgb=RGBColor.from_string('17365D')
sec.header.paragraphs[0].text='INFIMATCH  |  BACKEND V1  |  INVENTAIRE VAULT'
foot=sec.footer.paragraphs[0]
foot.text='Document sans valeurs secrètes • 15 septembre 2026   |   '
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)

def p(s): doc.add_paragraph(s)
def h(s): doc.add_heading(s,1)
def table(headers,rows):
 t=doc.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
 for c,s in zip(t.rows[0].cells,headers):c.text=s
 for row in rows:
  cells=t.add_row().cells
  for c,s in zip(cells,row):c.text=s
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
 return t
def bullet(s): doc.add_paragraph(s,style='List Bullet')

doc.add_heading('Ce que contient Vault',0)
p('InfiMatch • Inventaire de l’installation locale V1 • 15 septembre 2026')
p('Vault centralise les secrets nécessaires au backend et à son infrastructure. Ce document décrit leurs noms, leurs usages et les accès autorisés. Il ne contient aucune valeur de mot de passe, de clé ou de jeton.')
h('1. État constaté')
proof=json.loads((root/'docs/proofs/vault-env-comparison.json').read_text(encoding='utf-8'))
p('Dernière comparaison enregistrée : '+proof['checkedAt']+' (UTC). Vault était initialisé et déverrouillé. Les deux ensembles de secrets étaient en version 1, sans manque, différence ou clé supplémentaire par rapport au .env.')
table(['Élément','Contenu'],[
('Moteur de secrets','KV v2 : stockage de paires nom/valeur avec versions.'),
('Dossier backend','kv/infimatch/v1/backend : 8 entrées.'),
('Dossier infrastructure','kv/infimatch/v1/infra : 4 entrées.'),
('Total','12 entrées réparties dans deux dossiers ; 11 secrets distincts car SERVICE_TOKEN est partagé.'),
('Adresse locale','https://127.0.0.1:58200 — publiée uniquement sur la boucle locale du poste.'),
('Persistance','Stockage Raft dans un volume Docker ; journal d’audit dans un autre volume.'),
])
h('2. Ce que Vault ne remplace pas')
p('Les profils, missions et candidatures restent dans PostgreSQL. Les résultats et explications du matching interne restent dans MongoDB. Les documents privés restent dans le stockage de fichiers chiffrés. Vault conserve les secrets permettant d’utiliser ces services, pas leurs données métier.')
p('Le .env historique existe encore pour les anciennes commandes. L’installation locale de Vault ne prouve pas à elle seule la sécurisation complète de l’environnement livré.')

doc.add_page_break()
h('3. Secrets du backend')
p('Chemin logique : kv/infimatch/v1/backend. Chemin API KV v2 : kv/data/infimatch/v1/backend.')
table(['Nom','Utilité'],[
('DATABASE_URL','Chaîne de connexion à PostgreSQL, contenant notamment les identifiants de connexion.'),
('MONGODB_URI','Chaîne de connexion à MongoDB pour les résultats et explications du matching.'),
('SESSION_SECRET','Secret de sécurité des sessions du backend.'),
('DOCUMENT_KEY','Clé de chiffrement des documents privés. Sa conservation est nécessaire pour relire les fichiers concernés.'),
('SERVICE_TOKEN','Secret partagé pour authentifier les échanges internes avec les automatisations n8n.'),
('RPPS_API_KEY','Clé d’accès au fournisseur Annuaire Santé / RPPS.'),
('FT_CLIENT_ID','Identifiant client de l’intégration France Travail.'),
('FT_CLIENT_SECRET','Secret client associé à l’intégration France Travail.'),
])
p('Le chargeur prévoit également les noms DOCUMENT_KEY_Vn pour conserver des clés documentaires versionnées. Aucun de ces champs supplémentaires ne figure dans l’inventaire vérifié.')
h('4. Secrets de l’infrastructure')
p('Chemin logique : kv/infimatch/v1/infra. Chemin API KV v2 : kv/data/infimatch/v1/infra.')
table(['Nom','Utilité'],[
('POSTGRES_PASSWORD','Mot de passe fourni à la configuration du service PostgreSQL.'),
('MONGO_PASSWORD','Mot de passe fourni à la configuration du service MongoDB.'),
('N8N_ENCRYPTION_KEY','Clé utilisée par n8n pour chiffrer ses identifiants enregistrés.'),
('SERVICE_TOKEN','Même secret de communication interne que dans le dossier backend.'),
])
p('Copier une nouvelle valeur dans Vault ne modifie pas automatiquement le mot de passe d’une base déjà initialisée. Tout changement doit rester cohérent entre le serveur et ses clients.')

doc.add_page_break()
h('5. Qui peut lire ou modifier ces secrets ?')
p('L’accès applicatif utilise AppRole : une identité de rôle et un identifiant secret permettent d’obtenir un jeton temporaire.')
table(['Rôle','Permissions configurées'],[
('infimatch-v1-backend','Lecture du seul dossier backend V1. Pas d’écriture, ni d’accès au dossier infra, à V2 ou à l’administration.'),
('infimatch-v1-infra','Lecture du seul dossier infra V1. Pas d’accès aux secrets backend.'),
('infimatch-v1-operator','Lecture et écriture des secrets V1, lecture des métadonnées, gestion des SecretID AppRole V1 et création de snapshots Raft.'),
])
p('Les rôles ne reçoivent pas la policy default. Les jetons ont une durée configurée de 5 minutes, avec un maximum de 10 minutes. Les lanceurs révoquent leur jeton après lecture, avant de démarrer le processus.')
p('Durées configurées des SecretID : 7 jours pour backend et infra, 30 jours pour operator. Le renouvellement doit être organisé avant expiration. Le jeton root initial a été révoqué ; aucun jeton root permanent n’est prévu pour l’usage courant.')
h('6. Fichiers associés, conservés hors du KV')
p('Le répertoire local data/vault/ est ignoré par Git et protégé par les droits du poste. Il contient des moyens d’accès ou de récupération sensibles, distincts des 11 secrets métier inventoriés.')
table(['Fichiers / éléments','Rôle'],[
('backend.json, infra.json, operator.json','Identifiants AppRole locaux, accessors et informations de durée de vie.'),
('recovery.json','Parts de déverrouillage Shamir : trois parts générées, deux nécessaires.'),
('ca.key et tls/','Clés privées et certificats utilisés pour HTTPS local.'),
('runtime.json','Configuration d’exécution non secrète : environnement, port, origine et paramètres autorisés.'),
('raft-*.snap','Copies de sauvegarde du stockage Raft, à traiter comme sensibles.'),
])
p('Les trois parts sont actuellement regroupées sur le poste. Leur séparation entre responsables et leur sauvegarde sécurisée hors du poste restent à organiser. Ces fichiers ne doivent pas être inclus dans les archives de conversation.')

doc.add_page_break()
h('7. Utilisation et mise à jour')
table(['Commande depuis InfiMatch','Effet'],[
('npm run vault:status','Affiche l’état de Vault, sans afficher les secrets.'),
('npm run vault:check','Vérifie le chargement de la configuration et la connexion PostgreSQL.'),
('npm run start:vault','Démarre l’API compilée avec les secrets lus dans Vault.'),
('npm run worker:vault','Démarre le worker avec les secrets lus dans Vault.'),
('npm run vault:sync','Copie les valeurs autorisées du .env vers le KV et met à jour runtime.json.'),
('npm run vault:snapshot','Crée un snapshot privé du stockage Raft.'),
])
p('La synchronisation est explicite : aucun suivi permanent du .env n’est installé. Avant toute synchronisation, comparer les valeurs en mémoire, comprendre les écarts et ne transférer que les changements voulus. Ne jamais afficher les valeurs dans les logs ou dans ce document.')
p('Après modification d’un secret utilisé par un processus, prévoir son redémarrage contrôlé et vérifier le fonctionnement. Les lanceurs lisent les secrets au démarrage ; ils ne les actualisent pas en continu. En mode Vault, le backend ne se rabat pas sur le .env si Vault est indisponible.')
p('La commande vault:rotate prépare le renouvellement des accès AppRole ; sa rotation réelle n’a pas été validée. Elle est distincte de la synchronisation du .env et de la rotation des mots de passe des bases ou des clés documentaires.')
h('8. Vérifications et limites')
for s in [
'Preuves précédentes : 10 tests Vault réussis ; reprise après redémarrage, persistance, connexion PostgreSQL et santé d’une API de recette vérifiées.',
'La comparaison du 15 septembre confirme 11 secrets distincts à jour ; aucune écriture supplémentaire n’était nécessaire.',
'Restent à valider : restauration sur une instance distincte, renouvellement réel des accès et bascule des processus utilisant encore le .env.',
'Interface Edge corrigée le 15 septembre 2026 : formulaire de connexion affiché, HTTPS et AppRole conservés ; 10 tests Vault réussis.',
'Le dossier V2 existe pour la préparation du projet. Aucun dossier de secrets V2 n’est inclus dans cet inventaire.'
]:bullet(s)
h('Sources locales')
p('scripts/vault/common.mjs ; scripts/vault/manage.mjs ; scripts/vault/run.mjs ; infra/vault/server.hcl ; infra/vault/compose.yaml ; docs/VAULT_V1.md ; docs/proofs/vault-env-comparison.json et vault-lifecycle.json.')
out=root/'docs/CONTENU_VAULT_V1.docx'
doc.save(out)
check=Document(out)
text='\n'.join([p.text for p in check.paragraphs]+[c.text for t in check.tables for r in t.rows for c in r.cells])
for k in ['DATABASE_URL','MONGODB_URI','SESSION_SECRET','DOCUMENT_KEY','SERVICE_TOKEN','RPPS_API_KEY','FT_CLIENT_ID','FT_CLIENT_SECRET','POSTGRES_PASSWORD','MONGO_PASSWORD','N8N_ENCRYPTION_KEY']:
 assert k in text,k
assert len(check.tables)==6
stamp=datetime.now().strftime('%Y%m%d-%H%M%S')
history=root/'docs/history/conversations'/f'{stamp}-document-contenu-vault.md'
history.write_text('# Document du contenu de Vault\n\nDemande : produire un document décrivant le contenu de Vault.\n\nLivré : docs/CONTENU_VAULT_V1.docx. Inventaire des 11 secrets distincts, usages, permissions, fichiers privés hors KV, commandes et limites. Aucune valeur secrète incluse. Source : configuration locale et dernière comparaison enregistrée. Aucun secret modifié.\n',encoding='utf-8')
archive=root.parent/'backups'/f'InfiMatch_document_Vault_{stamp}.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
 for f in [out,history]:z.write(f,f.relative_to(root))
with zipfile.ZipFile(archive) as z:assert z.testzip() is None
print(str(out))
print('DOCX relu : 11 noms présents, 6 tableaux ; archive vérifiée.')
