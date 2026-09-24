"""Compare proposed prices with disclosed costs, not net company profit."""
import csv
import json
import math
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent
v = json.loads((ROOT.parents[2]/'annexe/rendu/2026-09-23/resultats-couts.json').read_text(encoding='utf8'))
COST = v['monthly_total']
CLIENTS = 20
CAP = 30
prices = [('mission', 'À la mission — 19 € par validation', lambda n: 19*n, 1000),
          ('forfait', 'Forfait — 299 €/mois × 20 établissements', lambda n: 299*CLIENTS, CLIENTS*CAP),
          ('mixte', 'Mixte — 149 €/mois × 20 établissements + 9 €/validation', lambda n: 149*CLIENTS+9*n, 1000)]
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False})
euro = FuncFormatter(lambda x,pos: f'{x:,.0f} €'.replace(',',' '))
def style(ax):
    ax.set_facecolor('#f4f6f8')
    ax.grid(axis='y',color='#dce3ea')
    ax.yaxis.set_major_formatter(euro)
    ax.set_xlabel('Validations mensuelles — total de la plateforme')

with PdfPages(ROOT/'rentabilite-trois-propositions.pdf') as pdf:
    for slug,title,revenue,limit in prices:
        fig,axes = plt.subplots(1,2,figsize=(14,7),facecolor='#f4f6f8')
        fig.suptitle(title,x=.07,ha='left',fontsize=19,fontweight='bold',color='#15283f')
        fig.text(.07,.91,'Scénario : 20 établissements payants • budget mensuel constant de 5 265,30 €',color='#536679')
        ns=list(range(limit+1)); sales=[revenue(n) for n in ns]; balances=[r-COST for r in sales]
        ax=axes[0];style(ax)
        ax.plot(ns,sales,label='Chiffre d’affaires HT',color='#234bff',lw=2.5)
        ax.axhline(COST,label='Budget de fonctionnement',color='#dc8739',ls='--')
        ax.plot(ns,balances,label='Solde après budget',color='#009d8f',lw=2.5)
        ax.axhline(0,color='#536679',lw=1)
        ax.set_title('Chiffre d’affaires, coûts et solde mensuels',loc='left',pad=15)
        ax.set_ylabel('Montant mensuel (€)')
        ax.legend(loc='upper left',fontsize=9)
        ax.set_xlim(0,1000)
        unit=axes[1];style(unit);us=list(range(50,limit+1))
        unit.plot(us,[revenue(n)/n for n in us],color='#234bff',lw=2,label='Recette moyenne / validation')
        unit.plot(us,[COST/n for n in us],color='#dc8739',ls='--',label='Coût moyen / validation')
        unit.plot(us,[(revenue(n)-COST)/n for n in us],color='#009d8f',lw=2.5,label='Solde moyen / validation')
        unit.axhline(0,color='#536679',lw=1)
        unit.set_title('Résultat moyen alloué par mission',loc='left',pad=15)
        unit.set_ylabel('Montant par validation (€)');unit.set_xlim(50,1000)
        unit.legend(loc='lower right',fontsize=9)
        if slug=='forfait':
            message='Couverture dès 18 abonnés. Avec 20 : +714,70 €/mois, indépendamment du nombre de validations.'
            for a in axes:
                a.axvspan(600,1000,color='#dce3ea',alpha=.65)
                a.text(.72,.48,'Au-delà de 600 :\noffre non définie\npour 20 abonnés',transform=a.transAxes,color='#536679',fontsize=10)
        else:
            threshold=math.ceil((COST-(CLIENTS*149 if slug=='mixte' else 0))/(9 if slug=='mixte' else 19))
            message=f'Couverture du budget à partir de {threshold} validations/mois'+(' avec 20 abonnés.' if slug=='mixte' else '.')
            ax.axvline(threshold,color='#536679',ls=':',lw=1)
            ax.annotate(f'{threshold} validations',(threshold,0),xytext=(12,-30),textcoords='offset points',fontweight='bold')
        fig.text(.07,.17,message,color='#15283f',fontsize=11,fontweight='bold')
        fig.text(.07,.09,'Solde du périmètre modélisé, pas bénéfice net : hors commercial, frais de paiement, fiscalité et dépassements.\nForfait : 30 validations maximum PAR établissement ; la répartition des missions doit respecter ce plafond.\nUne validation est une affectation confirmée, pas une prestation de soins réalisée. Coût unitaire indéfini à zéro.',fontsize=9,color='#536679')
        fig.subplots_adjust(left=.07,right=.98,top=.81,bottom=.29,wspace=.23)
        fig.savefig(ROOT/f'rentabilite-{slug}.png',dpi=150,facecolor=fig.get_facecolor())
        pdf.savefig(fig);plt.close(fig)

    fig,ax=plt.subplots(figsize=(12,6.8),facecolor='#f4f6f8');style(ax)
    for (slug,title,revenue,limit),color in zip(prices,['#234bff','#dc8739','#009d8f']):
        ns=list(range(50,limit+1));ax.plot(ns,[(revenue(n)-COST)/n for n in ns],label=title,color=color,lw=2.5)
    ax.axhline(0,color='#15283f',lw=1.5)
    ax.set(xlim=(50,1000),ylabel='Solde moyen par validation (€)')
    ax.legend(loc='lower right',fontsize=9)
    fig.suptitle('Quel solde par mission pour chaque proposition ?',x=.09,ha='left',fontsize=19,fontweight='bold',color='#15283f')
    fig.text(.09,.91,'20 établissements payants • coût = 5 265,30 €/mois • tarifs proposés, non commercialisés',color='#536679')
    fig.text(.09,.08,'Au-dessus de zéro : recettes supérieures au budget modélisé, avant les coûts exclus.\nForfait interrompu à 600 validations : 20 établissements × 30 maximum, sous réserve de leur répartition.\nPas d’extrapolation de coût au-delà des quotas. À zéro mission, le solde mensuel existe, pas le ratio par mission.',fontsize=9,color='#536679')
    fig.subplots_adjust(left=.09,right=.98,top=.84,bottom=.22)
    fig.savefig(ROOT/'rentabilite-comparaison.png',dpi=170,facecolor=fig.get_facecolor());pdf.savefig(fig);plt.close(fig)

volumes=[0,50,100,200,250,254,278,300,500,600,750,1000]
rows=[]
for n in volumes:
    for slug,title,revenue,limit in prices:
        valid=n<=limit
        rows.append([slug,n,CLIENTS,round(revenue(n),2) if valid else '',round(COST,2),round(revenue(n)-COST,2) if valid else '',round((revenue(n)-COST)/n,4) if valid and n else '', 'Applicable sous hypothèses' if valid else 'Hors forfait 20 établissements'])
with (ROOT/'rentabilite-scenarios.csv').open('w',encoding='utf8',newline='') as f:
    wr=csv.writer(f,delimiter=';');wr.writerow(['formule','validations','etablissements','ca_ht','cout','solde','solde_par_validation','statut']);wr.writerows(rows)

wb=Workbook();a=wb.active;a.title='Parametres'
for row in [('Paramètre','Valeur'),('Coût mensuel',COST),('Établissements payants',CLIENTS),('Prix par validation',19),('Forfait par établissement',299),('Validations par établissement incluses',30),('Abonnement mixte',149),('Prix variable mixte',9)]:a.append(row)
s=wb.create_sheet('Scenarios');s.append(['Validations','CA mission','Solde mission','Solde / mission','CA forfait','Solde forfait','Solde / mission','CA mixte','Solde mixte','Solde / mission'])
for i,n in enumerate(volumes,2):
    s.append([n,f'=A{i}*Parametres!$B$4',f'=B{i}-Parametres!$B$2',f'=IF(A{i}=0,"",C{i}/A{i})',f'=IF(A{i}>Parametres!$B$3*Parametres!$B$6,"Hors forfait",Parametres!$B$3*Parametres!$B$5)',f'=IF(ISNUMBER(E{i}),E{i}-Parametres!$B$2,"")',f'=IF(AND(A{i}>0,ISNUMBER(F{i})),F{i}/A{i},"")',f'=Parametres!$B$3*Parametres!$B$7+A{i}*Parametres!$B$8',f'=H{i}-Parametres!$B$2',f'=IF(A{i}=0,"",I{i}/A{i})'])
for sh in wb:
    sh.freeze_panes='B2'
    for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='15283F')
    for col in 'ABCDEFGHIJ':sh.column_dimensions[col].width=24
    for row in sh.iter_rows(min_row=2):
        for c in row:
            if c.data_type=='f' or isinstance(c.value,(int,float)):c.number_format='#,##0.00'
a.column_dimensions['A'].width=43
wb.save(ROOT/'rentabilite-scenarios.xlsx')
assert 19*277<COST<=19*278
assert 299*17<COST<=299*18
assert 20*149+9*253<COST<=20*149+9*254
assert len(rows)==36
print('PASS: 3 individual graphs, comparison, 4-page PDF, CSV, editable Excel and thresholds checked')
