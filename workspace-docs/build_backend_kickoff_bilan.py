from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import json, hashlib, zipfile, shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path('E:/Interimatch'); repo=root/'InfiMatch'; out=repo/'docs/BILAN_BACKEND_KICKOFF_V1.docx'
doc=Document(); sec=doc.sections[0]; sec.page_height=Cm(29.7); sec.page_width=Cm(21); sec.top_margin=Cm(1.65); sec.bottom_margin=Cm(1.6); sec.left_margin=sec.right_margin=Cm(1.7)
normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(10); normal.paragraph_format.space_after=Pt(5)
for n in ['Title','Heading 1','Heading 2']:
 doc.styles[n].font.color.rgb=RGBColor.from_string('173D55')
sec.header.paragraphs[0].text='INFIMATCH  /  BACKEND V1  /  BILAN DE REPRISE'
sec.header.paragraphs[0].runs[0].font.size=Pt(8)
f=sec.footer.paragraphs[0]; f.text='15 septembre 2026 • État documentaire et preuves existantes                                   '
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
def p(t,style=None): doc.add_paragraph(t,style)
def h(t): doc.add_heading(t,2)
def bullet(t): p(t,'List Bullet')
def page(t): doc.add_page_break(); doc.add_heading(t,1)
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1, cols=len(headers)); t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,headers): c.text=v
 repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row): c.text=v
  pr=t.rows[-1]._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
 for row in t.rows:
  for c in row.cells:
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(4)
    for r in para.runs: r.font.size=Pt(9)
 return t

doc.add_heading('Backend InfiMatch',0)
p('Bilan vis-à-vis du kick-off\nRéalisé • Reste à faire • Passage de relais','Subtitle')
p('État au 15 septembre 2026 | Équipe de 4 personnes | Projet de 11 jours')
h('1. Synthèse à partager')
p('Le backend fonctionne localement sur les parcours testés. Le socle technique, les parcours métier principaux, les intégrations France Travail / FINESS / Annuaire Santé et les trois automatisations locales sont présents. La V1 reste à finaliser et à valider dans son environnement de livraison avec le frontend.')
table(['Dernières preuves disponibles','Résultat / limite'],[
('Tests, typage et compilation','71 tests réussis ; typecheck et build réussis.'),
('Couverture générée','79,38 % des lignes ; 84,33 % des branches. Ce taux ne mesure pas le pourcentage de fonctionnalités terminées.'),
('Données publiques','134 annonces distinctes sur le lot Paris ; import et rejeu sans doublons. FINESS importé ; RPPS FOUND et NOT_FOUND vérifiés auprès du fournisseur.'),
('Automatisations','3 workflows n8n exécutés en local. Connexion Cloud non démontrée.'),
('Livraison','HTTPS, restauration complète, installation propre et recette frontend/backend restent à valider.')])
h('Priorités de clôture')
p('HTTPS et cadrage du contrôle réglementaire d’expérience → dépôts documentaires sans doublons et conservation → actualisation des offres → environnement n8n retenu et worker → restauration et recette finale.')
h('Comment lire ce bilan')
p('Implémenté : le code est présent. Testé : une preuve couvre des scénarios précis. À terminer : un développement manque. À valider : une preuve complète manque. Aucun de ces états ne constitue une conformité intégrale.')
p('Sources : sujet D-WEB-901-project.pdf, checklist fournie dans cette conversation, échange joint, historique du dépôt, documentation et preuves locales. Les documents joints sont des sources de cadrage ; leurs anciennes demandes ne sont pas exécutées comme de nouvelles instructions.')
p('La preuve de recette date du 14/09/2026 à 23:29 UTC, soit le 15/09 à 01:29 à Paris, et référence le code 5f9db50b26638a03259d37640d1aea8e24c53e18. Aucun test backend ni appel fournisseur n’a été relancé pour rédiger ce Word.')

page('2. Attentes du kick-off : couverture backend')
p('Les pages ci-dessous désignent les pages physiques du PDF, couverture comprise. Les choix InfiMatch complètent les attentes du sujet.')
table(['Attente du sujet','État constaté','À terminer / preuve de fin'],[
('Comptes entreprise / intérimaire, authentification sécurisée (p. 3)','Inscription, connexion, déconnexion, mots de passe hachés, sessions, Origin/CSRF ; entreprises séparées en agences et établissements.','Recette de toutes les routes sensibles, droits croisés et parcours navigateur.'),
('Données sensibles chiffrées au repos et en transit (p. 3)','Documents privés chiffrés au repos ; accès contrôlé ; tests d’altération et de rotation.','HTTPS de livraison, protection interservices, privilèges minimaux et gestion opérationnelle des clés.'),
('Missions, profils, matching et tableaux de bord (p. 3)','Champs métier, transitions, candidatures, affectation humaine, score interne expliqué ; données des tableaux de bord.','Recette exhaustive, contrôle d’expérience préalable applicable et intégration des écrans.'),
('Source publique nettoyée et utile dans le produit (p. 4)','France Travail : import réel, provenance, normalisation, dédoublonnage ; FINESS consultable ; comparaison partielle externe.','Rafraîchissement et retraits ; démontrer une utilisation visible dans le frontend et un exemple avant/après nettoyage.'),
('Au moins 2 automatisations nocode (p. 4)','3 workflows locaux : notification, relance, confirmation PDF. Exports JSON disponibles.','Rejouer les scénarios sur l’environnement livré ; raccorder Cloud seulement si retenu. Le sujet n’impose pas Cloud.'),
('Framework Node / TypeScript ; SQL et NoSQL (p. 6)','NestJS / TypeScript, PostgreSQL + PostGIS, MongoDB pour les explications de matching.','Installation reproductible et preuve des usages sur l’environnement livré.'),
('Tests unitaires, fonctionnels et couverture livrée (p. 6)','71 tests selon la dernière preuve ; rapport HTML et archive de couverture.','Compléter les cas manquants et produire un rapport associé à la révision finale.'),
('Bibliothèque CLI (p. 6)','Commander dans le backend ; commandes de migration et d’import/nettoyage.','Documenter et montrer une commande reproductible avec résultat et provenance.')])
p('Le kick-off exige un minimum de deux automatisations ; la V1 retenue en prévoit trois. Les suggestions Slack/Discord du sujet ne signifient pas qu’un tel canal est déjà branché. Les notifications actuelles sont internes à InfiMatch.')

page('3. Ce qui est réalisé dans la V1')
h('Socle et architecture')
p('Une application NestJS modulaire porte les règles métier. PostgreSQL est la base de référence ; PostGIS calcule les distances. MongoDB conserve les explications du matching interne. Les fichiers sont privés et chiffrés. Le worker traite les événements persistés et n8n orchestre les appels autorisés vers l’API.')
table(['Flux actuel','Fonction'],[
('Client → API /api/v1 → PostgreSQL','Comptes, profils, organisations, missions, candidatures, affectations, notifications et événements.'),
('API → PostGIS / MongoDB','Distance et admissibilité ; score interne explicable et versionné.'),
('API → France Travail / FINESS / ANS','Offres externes / référentiel établissements / contrôle RPPS.'),
('Événement SQL → worker → n8n → API','Notification, relance, génération de confirmation PDF et reçu métier final.'),
('API → stockage privé chiffré','Dépôt, téléchargement autorisé et confirmation PDF fictive.')])
h('Profils, missions et candidatures')
bullet('Profils IDE, IADE et IBODE : compétences, expérience déclarée, disponibilités, indisponibilités et mobilité. Vérification RPPS : retrouvé, introuvable ou en attente ; FINESS importé et consultable.')
bullet('Création, modification, publication, annulation, réouverture et clôture des missions ; candidature, retrait, sélection et refus. Affectation finale humaine par l’agence, avec contrôle des droits et conflits.')
bullet('Favoris, historique, données des tableaux de bord et notifications consultables / marquables comme lues. Principales commandes métier protégées contre les doublons au rejeu.')
h('Matching interne et annonces externes')
bullet('Interne : contrôles bloquants avant score, disponibilités complètes, qualifications, compétences obligatoires, conflits et mobilité ; classement des profils admissibles et règles expliquées/versionnées.')
bullet('Externe : comparaison privée au profil connecté ; distinction entre incompatibilité connue, donnée manquante et indice à confirmer ; option explicite pour inclure les annonces incomplètes. Aucun score global externe ni disponibilité inventée ; candidature par redirection.')
h('Acquisition réelle et documents')
bullet('France Travail : recherches croisées infirmier / IDE / IADE / IBODE. Lot Paris : 120 IDE, 7 IADE, 1 IBODE et 6 qualifications non confirmées, soit 134 offres distinctes ; rejeu sans doublons.')
bullet('FINESS : 174 621 identifiants, dont 104 752 actifs ; 120 663 entrées avec coordonnées. Snapshot daté du 1er septembre 2026. Une entrée FINESS ne confère aucun droit sur une organisation.')
bullet('Documents fictifs privés et confirmation PDF ; tests de chiffrement, altération, rotation et reprise. Code sauvegardé historiquement par GitHub / bundles ; restauration PostgreSQL isolée déjà exercée.')

page('4. Reste à faire : sécurité et fiabilité')
p('Responsables proposés, à répartir dans l’équipe : A = backend métier ; B = données et matching ; C = frontend ; D = infrastructure, documents et automatisations. Ce sont des périmètres de travail, pas des affectations nominatives confirmées.')
table(['Priorité / pilote','Travail à réaliser','Critère de validation'],[
('P1 • D + A\nTransport et accès','Configurer HTTPS ; protéger les échanges interservices ; contrôler les privilèges SQL/Mongo, clés et images déployées.','Parcours authentifié sur HTTPS ; services privés ; comptes techniques restreints ; configuration et vérifications consignées.'),
('P1 • A + B\nExpérience préalable','Cadrer la règle applicable puis implémenter équivalent temps plein (ETP), périodes retenues, profession/spécialité et justificatifs.','Règle sourcée et versionnée ; cas temps partiel, périodes chevauchantes, spécialités et justificatifs manquants testés ; résultat d’admissibilité explicite.'),
('P1 • A + D\nDocuments sans doublons','Étendre l’idempotence aux dépôts documentaires et remplacements bancaires, y compris après interruption.','Même commande rejouée : un seul résultat logique ; contenu divergent refusé ; concurrence et reprise testées ; droits revérifiés.'),
('P1 • D + A\nConservation','Définir les durées, purger les fichiers orphelins et compléter les mécanismes de conservation des données.','Politique documentée ; purge testée sans suppression de document encore lié ; anciennes clés conservées tant que nécessaires.'),
('P2 • A + B\nRoutes et contrats','Compléter tests de champs, transitions, permissions et pannes ; réponses complexes OpenAPI du matching, documents, données publiques et dashboards.','Cas nominaux et refus documentés ; schémas de sortie conformes aux réponses ; tests ciblés passants.')])
h('Point d’attention : expérience et RPPS')
p('Le code experienceMonths calcule une durée calendaire par service, en fusionnant les périodes. Ce calcul ne modélise pas à lui seul l’ETP ni le contrôle réglementaire demandé dans la checklist. Le RPPS retrouvé ne remplace pas ce contrôle, ni les autres conditions d’admissibilité.')
p('Ce bilan décrit le travail à effectuer : il ne fixe aucun seuil légal ni ne conclut sur son champ d’application. Les règles à implémenter doivent être vérifiées auprès des sources officielles applicables au scénario retenu avant développement.')
h('Point d’attention : automatisation')
p('Une affectation reste une décision humaine. La confirmation produite dans ce POC est un PDF fictif : elle ne prouve pas la conformité d’un contrat de travail et ne vaut pas signature électronique.')

page('5. Reste à faire : données et livraison')
table(['Priorité / pilote','Travail à réaliser','Critère de validation'],[
('P1 • B\nCycle des offres','Compléter le rafraîchissement et gérer les retraits confirmés du fournisseur.','Modification et retrait testés ; une absence dans une recherche partielle ne supprime pas une offre ; un échec fournisseur préserve le dernier état fiable.'),
('P1 • D\nn8n et worker','Raccorder n8n Cloud si retenu ; vérifier démarrage automatique, surveillance et reprise du worker.','3 scénarios exécutés sur la cible ; reçu SQL et résultat final vérifiés ; redémarrage et événement épuisé observables.'),
('P1 • D\nRelance','Corriger l’encodage du message de relance présent dans automation.module.ts.','Une nouvelle notification affiche correctement « Une mission reste à pourvoir. » ; définir le traitement des anciens messages si nécessaire.'),
('P2 • D + équipe\nRestauration','Restaurer ensemble PostgreSQL, MongoDB, fichiers, clés et n8n dans un environnement isolé.','Comptes et affectations cohérents ; documents lisibles avec les bonnes clés ; explications et workflows opérationnels ; procédure et résultat conservés.'),
('P2 • équipe\nInstallation propre','Repartir d’une copie propre du dernier état ; installer, migrer, démarrer API, worker et services.','README suffisant pour un autre membre ; aucune dépendance cachée au poste initial ; versions et configuration nécessaires documentées.'),
('P2 • A + B + C\nRecette et livraison','Intégrer le frontend, harmoniser les anciens statuts/chiffres, vérifier OpenAPI, générer les preuves finales.','Parcours critiques réalisés dans le navigateur ; révision testée identifiée ; couverture et exports prêts pour le rendu.')])
h('Améliorations métier à poursuivre')
bullet('Réévaluer les pondérations expérimentales du score interne avec les retours métier ; conserver leur version et leurs explications.')
bullet('Améliorer la reconnaissance des services et exigences des offres externes ; laisser les informations non interprétables inconnues ou à confirmer.')
h('Dépendance avec le frontend')
p('Cartes, formulaires, calendrier, tableaux de bord visuels, notifications affichées et résultats de matching sont du travail frontend. Le backend fournit leurs données. La recherche FINESS, les erreurs RPPS et les états des annonces incomplètes doivent être intégrés et testés ensemble.')
p('Contrat client : cookie de session, Origin et X-CSRF-Token pour les écritures ; renouveler le jeton après connexion. Les commandes métier documentées demandent Idempotency-Key : nouvelle clé pour une nouvelle action, même clé pour un rejeu réseau. Listes : limit/offset selon OpenAPI.')

page('6. Livrables collectifs et recette de fin')
h('Ce que le sujet attend au-delà du backend')
table(['Exigence / page PDF','Contribution attendue'],[
('Cadrage J+2, roadmap et chiffrage • p. 3 et 7','Plans et matrices présents ; confirmer la version retenue au go/nogo et le chiffrage. Compléter les temps réellement passés et les écarts ; ne pas inventer de charge restante.'),
('Étude de marché et pitch • p. 3, 7 et 8','Justifier le secteur santé, la concurrence, les besoins et la proposition de valeur. Préparer support et démonstration ; chaque membre participe.'),
('Données personnelles et cadre métier • p. 5','Documenter base légale, durées de conservation et mentions ; vérifier les règles du scénario intérim retenu. Relier les mécanismes backend aux engagements annoncés.'),
('Accessibilité, éco-conception, SEO • p. 5','Frontend : principes de base d’accessibilité, responsive, balises et sitemap. Équipe : au moins deux pratiques d’éco-conception réellement appliquées et documentées ; réflexion achat responsable/réemploi si pertinente.'),
('Dépôt et livrables techniques • p. 6 et 7','Frontend + backend, README, exports n8n, script de nettoyage, données et provenance, tests fonctionnels/unitaires et rapport de couverture.')])
h('Parcours de démonstration et preuves à conserver')
for t in [
'Inscription et connexion des deux catégories de comptes ; refus d’un accès non autorisé entre organisations.',
'Profil infirmier, disponibilités et RPPS : retrouvé, introuvable et fournisseur indisponible ; ne pas présenter une fixture comme un appel réel.',
'Création/publication d’une mission, matching expliqué, candidature et consentement, sélection puis affectation humaine ; conflit et rejeu sans doublon.',
'Notification de matching, relance d’une mission non pourvue, confirmation PDF privée ; contrôler le résultat métier après chaque workflow.',
'Recherche d’une annonce publique nettoyée et traçable ; comparer au profil, expliquer une information manquante et ouvrir la candidature externe.',
'Démarrage propre, preuve HTTPS et restauration ; exécuter la recette finale et joindre sa couverture à la révision livrée.']:
 bullet(t)
h('Condition pour annoncer la V1 terminée')
p('Fermer les écarts retenus dans ce bilan avec une preuve, terminer les parcours visibles du produit, livrer les documents du sujet et consigner explicitement les limites résiduelles. Le nombre de tests ne remplace pas cette vérification.')
p('Périmètre futur conservé : attestation et signature électronique V2 ; références hors V1 ; administration avancée, contrats et paie selon les versions futures prévues. Ce bilan ne les ajoute pas au chantier actuel.')

page('7. Passage de relais et sauvegarde de conversation')
h('Reprendre dans le bon ordre')
for t in [
'Ouvrir E:/Interimatch/InfiMatch. Lire git status avant toute modification ; conserver les changements existants.',
'Lire ce bilan, README.md, docs/REPRISE_BACKEND_V1.md, docs/RECETTE_BACKEND_V1.md et docs/REQUIREMENTS_V1.md. Pour les chiffres, privilégier docs/proofs/verification.json et coverage-totals.json.',
'Consulter docs/history/DISCUSSION.md et IMPLEMENTATION.md ; distinguer les états historiques du dernier état. Ne pas recréer ou réinitialiser les données pour reprendre.',
'Suivre le README pour démarrer les services, l’API compilée et le worker. API locale : 127.0.0.1:3100 ; Swagger : /api/docs ; n8n local : 127.0.0.1:55678.',
'Traiter les tâches P1, puis exécuter les vérifications adaptées. npm run verify lance typage, compilation et couverture ; les tests fonctionnels nécessitent les services locaux décrits dans le README.',
'Après chaque lot : actualiser exigences, preuves, historique et note de reprise ; sauvegarder les fichiers de travail. Après commit et arbre propre, npm run snapshot crée un bundle Git vérifié.']:
 bullet(t)
h('Sauvegarde effectuée pour cette demande')
p('L’échange joint est conservé tel que fourni dans docs/history/conversations. La demande actuelle, la checklist et les décisions utiles sont reprises sous forme de synthèse datée. Cette synthèse n’est pas un export intégral des conversations non accessibles.')
p('Une archive locale horodatée de ce Word et des documents de reprise accompagne cette livraison dans E:/Interimatch/backups ; un manifeste SHA-256 permet d’en vérifier le contenu. Les secrets, bases et fichiers privés sont exclus. Cette archive documentaire ne remplace pas une restauration complète de l’application.')
p('La sauvegarde est reprise à ce jalon de travail. Aucun service planifié de sauvegarde continue n’est installé par ce document ; aux prochains jalons, renouveler le journal et l’archive.')
h('Références pour vérifier les constats')
for t in [
'Sujet : E:/Interimatch/D-WEB-901-project.pdf. Échange joint : pasted-text.txt fourni par l’utilisateur ; checklist de cette conversation.',
'Preuves de tests : docs/proofs/verification.json, coverage-totals.json, test-and-coverage.txt ; code référencé : 5f9db50b26638a03259d37640d1aea8e24c53e18.',
'Fournisseurs : docs/proofs/france-travail-rectification.json, finess-live.json, ans-fhir-positive.json et external-partial-live.json. Automatisations : docs/proofs/n8n-executions.json.',
'Architecture / contrats : docs/SCHEMA_ARCHITECTURE_V1.md, docs/FLUX_V1.md, docs/OFFRES_EXTERNES_V1.md et docs/openapi.json.',
'Contrôles ciblés du code pour ce bilan : backend/src/domain/matching.ts (durée calendaire), backend/src/automation/automation.module.ts (encodage de relance), scripts/snapshot.mjs (portée du bundle).']:
 p(t)
doc.save(out)
# Save the source exchange and a clearly labeled summary, without claiming a full transcript.
stamp=datetime.now().strftime('%Y%m%d-%H%M%S'); conv=repo/'docs/history/conversations'; conv.mkdir(exist_ok=True)
shutil.copyfile(Path('C:/Users/lebre/.codex/attachments/9c0fb159-052f-4be3-b32f-71f92e718ebd/pasted-text.txt'),conv/f'{stamp}-echange-fourni.txt')
summary='''# Conversation et point de reprise — 15 septembre 2026

Synthèse des messages accessibles, pas un export intégral.

## Demandes actuelles
- Produire un Word complet sur le backend : attentes du kick-off, réalisé, reste à faire, passage de relais.
- Utiliser la checklist fournie, regarder l’historique et reprendre la sauvegarde régulière de conversation.
- Le PDF et les anciens échanges servent de sources ; aucune ancienne instruction n’est réexécutée par défaut.

## Checklist utilisateur reprise
- Socle réalisé : NestJS/TypeScript, PostgreSQL/PostGIS, MongoDB, comptes, sessions, Origin/CSRF, permissions et chiffrement privé. Restent HTTPS, échanges interservices, droits minimaux et recette exhaustive.
- Profils/missions : IDE/IADE/IBODE, RPPS, FINESS, transitions, candidatures, affectation humaine, favoris/historique/dashboards, idempotence métier. Restent tests complets et contrôle d’expérience réglementaire (ETP, périodes, profession/spécialité, justificatifs).
- Matching : admissibilité, score interne expliqué, comparaison partielle externe, inconnues et incompatibilités distinguées, option annonces incomplètes. Restent retours métier sur pondérations et amélioration de reconnaissance.
- France Travail : accès réel, requêtes croisées, normalisation/provenance, rejeu, alertes et routes. Restent actualisation/retraits sans assimiler recherche partielle à inventaire complet.
- n8n : notifications, trois workflows locaux, worker/outbox/reçus. Restent Cloud si retenu, démarrage/surveillance du worker et encodage de relance.
- Documents : fichiers fictifs privés, PDF, altération/rotation testées, sauvegardes historiques Git. Restent idempotence documentaire, orphelins, conservation et restauration SQL/Mongo/fichiers/clés/n8n.
- Livraison : 71 tests, 79,38 % lignes, Swagger/README/scripts/schémas. Restent réponses complexes, harmonisation documentaire, installation propre et recette frontend.
- Priorité demandée : HTTPS et cadre réglementaire, documents sans doublons, offres, n8n retenu, restauration/recette.
- Les affichages et écrans restent du frontend, avec recette commune API.

## Travail réalisé dans ce jalon
Lecture du sujet PDF, de l’échange joint, de l’historique et des documents de suivi ; recoupement des preuves JSON et contrôles ciblés du code. Création de docs/BILAN_BACKEND_KICKOFF_V1.docx. Les 71 tests sont une preuve antérieure, pas une nouvelle exécution. Aucun code backend, compte fournisseur ni déploiement modifié.
Archive documentaire locale horodatée avec manifeste SHA-256 prévue et contrôlée à la clôture. Sauvegarde par jalon, aucun planificateur permanent installé. Aucun secret ni base inclus.
'''
(conv/f'{stamp}-synthese-reprise.md').write_text(summary,encoding='utf-8')
for name in ['DISCUSSION.md','IMPLEMENTATION.md']:
 path=repo/'docs/history'/name
 with path.open('a',encoding='utf-8') as f:
  f.write('\n\n## 2026-09-15 — Bilan Word kick-off et sauvegarde de reprise\nDemande : Word réalisé/reste à faire, intégration de la checklist utilisateur, lecture historique et reprise des sauvegardes de conversation. Livrable : docs/BILAN_BACKEND_KICKOFF_V1.docx. Sources recoupées : sujet, historique, preuves et code ciblé. 71 tests / 79,38 % lignes / 84,33 % branches sont les dernières preuves existantes, sans nouvelle exécution. Synthèse et échange joint archivés dans docs/history/conversations. Archive documentaire horodatée et manifeste SHA-256 dans E:/Interimatch/backups ; exclut secrets, bases et fichiers privés. Pas de planificateur permanent installé.\n')
# Check the Word XML and preserve a verifiable, narrowly scoped archive.
reopened=Document(out); texts=[p.text for p in reopened.paragraphs]+[c.text for t in reopened.tables for r in t.rows for c in r.cells]
assert len(reopened.tables)==6
assert all(s in '\n'.join(texts) for s in ['71 tests','79,38','ETP','kick-off','Idempotency-Key'])
assert not any('\ufffd' in s for s in texts)
files=[out,repo/'README.md',repo/'docs/REPRISE_BACKEND_V1.md',repo/'docs/REQUIREMENTS_V1.md',repo/'docs/RECETTE_BACKEND_V1.md',repo/'docs/history/DISCUSSION.md',repo/'docs/history/IMPLEMENTATION.md',repo/'docs/proofs/verification.json',repo/'docs/proofs/coverage-totals.json']+list(conv.glob(f'{stamp}-*'))
manifest={'createdAt':datetime.now().isoformat(),'scope':'Documentary checkpoint only; excludes secrets, databases and private files','files':[{ 'path':str(f.relative_to(repo)).replace('\\','/'),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in files]}
backup=root/'backups'/f'InfiMatch_Bilan_Kickoff_{stamp}.zip'
with zipfile.ZipFile(backup,'w',zipfile.ZIP_DEFLATED) as z:
 for f in files: z.write(f,f.relative_to(repo))
 z.writestr('MANIFEST_SHA256.json',json.dumps(manifest,ensure_ascii=False,indent=2))
with zipfile.ZipFile(backup) as z:
 assert z.testzip() is None
 for item in manifest['files']: assert hashlib.sha256(z.read(item['path'])).hexdigest()==item['sha256']
print(json.dumps({'word':str(out),'bytes':out.stat().st_size,'tables':len(reopened.tables),'backup':str(backup),'archivedFiles':len(files),'archiveVerified':True},ensure_ascii=False))
